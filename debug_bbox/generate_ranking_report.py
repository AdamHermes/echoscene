import os
import glob
import re
import json
import math
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from shapely.geometry import Polygon, MultiPolygon, Point
from shapely.ops import unary_union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenPyxlImage

# Load 3D-FRONT relationships for dynamic relational evaluation fallback
BASE_DIR = "/Users/lehoangan/Documents/GitHub/ROOM/echoscene"
REL_DATA_ALL = {}
for rf in ['FRONT/relationships_bedroom_test.json', 'FRONT/relationships_diningroom_test.json', 'FRONT/relationships_all_test.json']:
    p = os.path.join(BASE_DIR, rf)
    if os.path.exists(p):
        with open(p, 'r') as f:
            data = json.load(f)
            for scan in data.get('scans', []):
                REL_DATA_ALL[scan['scan']] = scan

def natural_sort_key(s):
    """Sorts strings using natural numerical ordering (matching VS Code order)."""
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', s)]

def detect_target_scene(debug_dir):
    """Auto-detects the target scene ID from txt filenames or log content."""
    for fpath in glob.glob(os.path.join(debug_dir, "*.txt")):
        fname = os.path.basename(fpath)
        m = re.search(r'debug_bbox_input_([A-Za-z0-9\-]+)\.txt', fname)
        if m:
            return m.group(1)
        with open(fpath, 'r') as f:
            for line in f:
                m2 = re.search(r'SCENE:\s*([A-Za-z0-9\-]+)', line)
                if m2:
                    return m2.group(1)
    return None

def get_obb_polygon(x, z, l, w, angle_deg):
    """Calculates 2D Shapely Polygon for Oriented Bounding Box."""
    angle_rad = np.radians(angle_deg)
    cos_a, sin_a = np.cos(angle_rad), np.sin(angle_rad)
    corners = np.array([[-l/2, -w/2], [l/2, -w/2], [l/2, w/2], [-l/2, w/2]])
    rotated = []
    for cx, cz in corners:
        rx = cx * cos_a - cz * sin_a
        rz = cx * sin_a + cz * cos_a
        rotated.append((x + rx, z + rz))
    return Polygon(rotated)

def parse_log_file(file_path):
    """Parses raw object bounding box details, pairwise overlaps, and explicit relational metrics from a log file."""
    objects = []
    raw_overlaps = []
    rel_acc_log = None
    means_of_mean_log = None
    
    if not os.path.exists(file_path):
        return objects, raw_overlaps, rel_acc_log, means_of_mean_log
        
    with open(file_path, 'r') as f:
        lines = f.readlines()
        
    content = "".join(lines)
    
    # 1. Look for explicit evaluation lines: acc & ... Total: &0.98 or Total: 0.98
    acc_m = re.search(r'acc &.*Total:\s*&?\s*([\d\.]+)', content)
    if acc_m:
        val = float(acc_m.group(1))
        rel_acc_log = val * 100.0 if val <= 1.0 else val
        
    means_m = re.search(r'means of mean:\s*([\d\.]+)', content)
    if means_m:
        val = float(means_m.group(1))
        means_of_mean_log = val * 100.0 if val <= 1.0 else val

    start_parsing = False
    for line in lines:
        if "---" in line:
            start_parsing = True
            continue
        if "Pairwise overlap" in line or "==" in line:
            start_parsing = False
        if "!! OVERLAP:" in line:
            raw_overlaps.append(line.strip())
            
        if start_parsing and line.strip():
            parts = line.split()
            if len(parts) >= 8:
                name = parts[0]
                try:
                    l = float(parts[1])
                    h = float(parts[2])
                    w = float(parts[3])
                    x = float(parts[4])
                    y = float(parts[5])
                    z = float(parts[6])
                    angle = float(parts[7].replace('°', ''))
                    objects.append({
                        "name": name, "l": l, "h": h, "w": w,
                        "x": x, "y": y, "z": z, "angle": angle
                    })
                except ValueError:
                    continue
    return objects, raw_overlaps, rel_acc_log, means_of_mean_log

def check_file_relations(objects, target_scene):
    """Evaluates individual GT spatial relationships for an object layout."""
    if not target_scene or target_scene not in REL_DATA_ALL:
        return []
        
    scan_rel = REL_DATA_ALL[target_scene]
    rel_objects = scan_rel['objects']
    relationships = scan_rel['relationships']
    
    furn_objs = [o for o in objects if o['name'] not in ['_scene_']]
    parsed_by_idx = {}
    keys_sorted = list(rel_objects.keys())
    for idx, key in enumerate(keys_sorted):
        if idx < len(furn_objs):
            parsed_by_idx[str(key)] = furn_objs[idx]
            
    results = []
    for rel in relationships:
        src, dest, rel_idx, rel_type = rel
        src_str, dest_str = str(src), str(dest)
        
        if src_str not in parsed_by_idx or dest_str not in parsed_by_idx:
            continue
            
        obj_s = parsed_by_idx[src_str]
        obj_d = parsed_by_idx[dest_str]
        
        if obj_s['name'] == 'floor' or obj_d['name'] == 'floor':
            continue
            
        passed = False
        val_str = ''
        
        if rel_type == 'left':
            diff = obj_s['z'] - obj_d['z']
            passed = diff < 0.05
            val_str = f'dz={diff:+.3f}m'
        elif rel_type == 'right':
            diff = obj_s['z'] - obj_d['z']
            passed = diff > -0.05
            val_str = f'dz={diff:+.3f}m'
        elif rel_type == 'front':
            diff = obj_s['x'] - obj_d['x']
            passed = diff > -0.05
            val_str = f'dx={diff:+.3f}m'
        elif rel_type == 'behind':
            diff = obj_s['x'] - obj_d['x']
            passed = diff < 0.05
            val_str = f'dx={diff:+.3f}m'
        elif rel_type in ['above', 'standing on']:
            diff = obj_s['y'] - obj_d['y']
            passed = diff >= -0.1
            val_str = f'dy={diff:+.3f}m'
        elif rel_type == 'close by':
            dist = math.sqrt((obj_s['x']-obj_d['x'])**2 + (obj_s['z']-obj_d['z'])**2)
            passed = dist < 2.5
            val_str = f'dist={dist:.3f}m'
        elif rel_type == 'bigger than':
            vol_s = obj_s['l'] * obj_s['h'] * obj_s['w']
            vol_d = obj_d['l'] * obj_d['h'] * obj_d['w']
            passed = vol_s >= vol_d * 0.8
            val_str = f'vol_s={vol_s:.2f}m³ vs vol_d={vol_d:.2f}m³'
        elif rel_type == 'smaller than':
            vol_s = obj_s['l'] * obj_s['h'] * obj_s['w']
            vol_d = obj_d['l'] * obj_d['h'] * obj_d['w']
            passed = vol_s <= vol_d * 1.2
            val_str = f'vol_s={vol_s:.2f}m³ vs vol_d={vol_d:.2f}m³'
        else:
            passed = True
            val_str = 'ok'
            
        results.append({
            'src_name': obj_s['name'],
            'src_id': src_str,
            'dest_name': obj_d['name'],
            'dest_id': dest_str,
            'relation': rel_type,
            'passed': passed,
            'detail': val_str
        })
    return results

def analyze_scene(objects, raw_overlaps, rel_acc_val, rel_details):
    """Computes layout quality metrics (collisions, OBB overlap area, out-of-bounds, walkability, relational acc)."""
    floor_obj = None
    furniture_objs = []
    for obj in objects:
        if obj["name"] == "floor":
            floor_obj = obj
        elif obj["name"] not in ["_scene_", "lamp"]:
            furniture_objs.append(obj)
            
    if floor_obj is not None:
        floor_poly = get_obb_polygon(floor_obj["x"], floor_obj["z"], floor_obj["l"], floor_obj["w"], floor_obj["angle"])
    else:
        floor_poly = Polygon([[-2.0, -2.5], [2.0, -2.5], [2.0, 2.5], [-2.0, 2.5]])
        
    furn_polys = []
    for obj in furniture_objs:
        poly = get_obb_polygon(obj["x"], obj["z"], obj["l"], obj["w"], obj["angle"])
        furn_polys.append((obj, poly))
        
    furn_aabb_overlaps = 0
    for line in raw_overlaps:
        match = re.search(r'OVERLAP:\s*(\w+)\s*<->\s*(\w+)', line)
        if match:
            o1, o2 = match.group(1), match.group(2)
            if o1 not in ["floor", "lamp", "_scene_"] and o2 not in ["floor", "lamp", "_scene_"]:
                furn_aabb_overlaps += 1
                
    obb_overlap_count = 0
    total_obb_overlap_area = 0.0
    N = len(furn_polys)
    for i in range(N):
        for j in range(i+1, N):
            p1 = furn_polys[i][1]
            p2 = furn_polys[j][1]
            if p1.intersects(p2):
                inter = p1.intersection(p2)
                if inter.area > 1e-4:
                    obb_overlap_count += 1
                    total_obb_overlap_area += inter.area
                    
    oob_count = 0
    max_oob_dist = 0.0
    total_oob_area = 0.0
    
    for obj, poly in furn_polys:
        if not floor_poly.contains(poly):
            diff = poly.difference(floor_poly)
            if diff.area > 1e-4:
                oob_count += 1
                total_oob_area += diff.area
                coords = np.array(poly.exterior.coords)
                for cx, cz in coords:
                    pt_dist = floor_poly.distance(Point(cx, cz))
                    if pt_dist > max_oob_dist:
                        max_oob_dist = pt_dist

    if furn_polys:
        all_furn_union = unary_union([p for _, p in furn_polys])
        free_floor = floor_poly.difference(all_furn_union)
    else:
        free_floor = floor_poly
        
    free_floor_area = free_floor.area
    free_floor_pct = (free_floor_area / max(floor_poly.area, 1e-5)) * 100.0
    
    if isinstance(free_floor, Polygon):
        walkable_components = 1 if free_floor.area > 0.01 else 0
    elif isinstance(free_floor, MultiPolygon):
        walkable_components = sum(1 for p in free_floor.geoms if p.area > 0.05)
    else:
        walkable_components = 0
        
    center_pt = Point(0, 0)
    min_center_dist = 999.0
    for _, poly in furn_polys:
        d = poly.distance(center_pt)
        if d < min_center_dist:
            min_center_dist = d
    if min_center_dist == 999.0:
        min_center_dist = 0.0

    # Counts of RIGHT and WRONG relations
    right_count = sum(1 for r in rel_details if r['passed'])
    wrong_count = sum(1 for r in rel_details if not r['passed'])
    wrong_summary = ", ".join([f"{r['src_name']}({r['src_id']})--[{r['relation']}]-->{r['dest_name']}({r['dest_id']})" for r in rel_details if not r['passed']])
    if not wrong_summary:
        wrong_summary = "None (All Relations RIGHT)"

    # Overall Quality Score
    score = 100.0
    score -= furn_aabb_overlaps * 20.0
    score -= total_obb_overlap_area * 35.0
    score -= oob_count * 15.0
    score -= max_oob_dist * 25.0
    score -= max(0, walkable_components - 1) * 8.0
    score += min(10.0, free_floor_pct * 0.15)
    score = float(np.clip(score, 0.0, 100.0))
    
    if score >= 90.0:
        status = "EXCELLENT"
    elif score >= 75.0:
        status = "GOOD"
    elif score >= 50.0:
        status = "MINOR ISSUES"
    else:
        status = "SEVERE COLLISIONS"
        
    return {
        "score": score,
        "status": status,
        "relational_acc": round(rel_acc_val, 1),
        "right_count": right_count,
        "wrong_count": wrong_count,
        "wrong_summary": wrong_summary,
        "rel_details": rel_details,
        "furn_aabb_overlaps": furn_aabb_overlaps,
        "obb_overlap_count": obb_overlap_count,
        "total_obb_overlap_area": round(total_obb_overlap_area, 4),
        "oob_count": oob_count,
        "max_oob_dist": round(max_oob_dist, 4),
        "free_floor_pct": round(free_floor_pct, 1),
        "walkable_components": walkable_components,
        "center_clearance": round(min_center_dist, 3)
    }

def render_scene_image(file_path, objects, metrics, out_img_path):
    """Renders high-quality 2D top-down layout plot image with metric overlay."""
    fig, ax = plt.subplots(figsize=(5.5, 5.5), dpi=150)
    
    color_map = {
        "bed": "#3498db", "chair": "#e74c3c", "nightstand": "#9b59b6", 
        "table": "#e67e22", "wardrobe": "#2ecc71", "lamp": "#f1c40f", 
        "floor": "#ecf0f1"
    }
    
    floor_obj = next((o for o in objects if o["name"] == "floor"), None)
    if floor_obj:
        f_poly = get_obb_polygon(floor_obj["x"], floor_obj["z"], floor_obj["l"], floor_obj["w"], floor_obj["angle"])
        f_patch = patches.Polygon(np.array(f_poly.exterior.coords), closed=True, facecolor="#f8f9fa", edgecolor="#7f8c8d", linewidth=2.0)
        ax.add_patch(f_patch)
        
    for obj in objects:
        if obj["name"] in ["_scene_", "floor"]:
            continue
        poly = get_obb_polygon(obj["x"], obj["z"], obj["l"], obj["w"], obj["angle"])
        c = color_map.get(obj["name"], "#1abc9c")
        alpha = 0.7 if obj["name"] != "lamp" else 0.4
        
        patch = patches.Polygon(np.array(poly.exterior.coords), closed=True, facecolor=c, edgecolor="#2c3e50", alpha=alpha, linewidth=1.5)
        ax.add_patch(patch)
        
        if obj["name"] != "lamp":
            ax.text(obj["x"], obj["z"], obj["name"], ha='center', va='center', fontsize=8, weight='bold',
                    bbox=dict(facecolor='white', alpha=0.75, edgecolor='none', pad=1.5))
                    
    ax.set_aspect('equal')
    ax.set_xlim(-3.2, 3.2)
    ax.set_ylim(-3.2, 3.2)
    ax.grid(True, linestyle=':', alpha=0.4)
    
    score = metrics["score"]
    rel_acc = metrics["relational_acc"]
    status = metrics["status"]
    
    fname = os.path.basename(file_path)
    if len(fname) > 40:
        fname = fname[:37] + "..."
    ax.set_title(f"{fname}\nScore: {score:.1f}/100 [{status}] | Rel Acc: {rel_acc:.1f}%", fontsize=8.5, weight='bold', color='black', pad=6)
    
    plt.tight_layout()
    os.makedirs(os.path.dirname(out_img_path), exist_ok=True)
    plt.savefig(out_img_path, dpi=150, bbox_inches='tight')
    plt.close(fig)

def build_excel_report(ranked_data, output_excel):
    """Creates a formatted Excel report with embedded layout images, summary table, and detailed relation breakdown sheet."""
    wb = openpyxl.Workbook()
    
    # ==========================================
    # SHEET 1: Layout Rankings Summary
    # ==========================================
    ws1 = wb.active
    ws1.title = "Layout Rankings Summary"
    ws1.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    sub_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
    
    font_title = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_red = Font(name="Segoe UI", size=10, bold=True, color="9C0006")
    font_green = Font(name="Segoe UI", size=10, bold=True, color="006100")
    
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    ws1.merge_cells("A1:P2")
    cell_title = ws1["A1"]
    cell_title.value = "  ROOM Layout Quality & Relational Accuracy Benchmark Report"
    cell_title.font = font_title
    cell_title.fill = header_fill
    cell_title.alignment = Alignment(vertical="center", horizontal="left")
    
    ws1.merge_cells("A3:P3")
    cell_sub = ws1["A3"]
    cell_sub.value = f"  Total Runs Analyzed: {len(ranked_data)}  |  Metrics: Relational Accuracy, RIGHT/WRONG Relations, Furniture Overlaps, OBB Overlap Area, Out-of-Bounds, Free Floor %, Walkability"
    cell_sub.font = Font(name="Segoe UI", size=10, italic=True, color="1F4E78")
    cell_sub.fill = sub_header_fill
    cell_sub.alignment = Alignment(vertical="center", horizontal="left")
    
    headers1 = [
        "Rank", "Layout Visualization", "Log File Name", "Quality Score", "Relational Acc (%)",
        "RIGHT Relations", "WRONG Relations", "WRONG Relation Details",
        "Overlaps (Count)", "OBB Overlap (m²)", "Out of Bounds", "Max OOB Dist (m)",
        "Free Floor %", "Walkable Components", "Center Clear (m)", "Grade Status"
    ]
    
    ws1.row_dimensions[5].height = 28
    for col_num, h_text in enumerate(headers1, 1):
        cell = ws1.cell(row=5, column=col_num)
        cell.value = h_text
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    col_widths1 = {
        "A": 8,   "B": 24,  "C": 45,  "D": 15,  "E": 18,  "F": 16,
        "G": 16,  "H": 40,  "I": 16,  "J": 18,  "K": 15,  "L": 18,
        "M": 14,  "N": 20,  "O": 16,  "P": 20
    }
    for col_letter, width in col_widths1.items():
        ws1.column_dimensions[col_letter].width = width

    start_row = 6
    for idx, item in enumerate(ranked_data, 1):
        row_num = start_row + idx - 1
        ws1.row_dimensions[row_num].height = 115
        
        m = item["metrics"]
        fname = os.path.basename(item["file_path"])
        img_path = item["img_path"]
        
        cell_rank = ws1.cell(row=row_num, column=1, value=f"#{idx}")
        cell_rank.alignment = Alignment(horizontal="center", vertical="center")
        cell_rank.font = font_bold
        
        if os.path.exists(img_path):
            img = OpenPyxlImage(img_path)
            img.width = 140
            img.height = 140
            ws1.add_image(img, f"B{row_num}")
            
        cell_file = ws1.cell(row=row_num, column=3, value=fname)
        cell_file.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell_file.font = font_body
        
        # Quality Score
        score_val = m["score"]
        cell_score = ws1.cell(row=row_num, column=4, value=score_val)
        cell_score.alignment = Alignment(horizontal="center", vertical="center")
        cell_score.font = Font(name="Segoe UI", size=11, bold=True)
        cell_score.number_format = '0.0'
        
        if score_val >= 90:
            cell_score.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        elif score_val >= 75:
            cell_score.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        elif score_val >= 50:
            cell_score.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        else:
            cell_score.fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
            
        # Relational Accuracy (%) Column
        rel_acc_val = m["relational_acc"]
        cell_rel = ws1.cell(row=row_num, column=5, value=rel_acc_val)
        cell_rel.alignment = Alignment(horizontal="center", vertical="center")
        cell_rel.font = Font(name="Segoe UI", size=11, bold=True)
        cell_rel.number_format = '0.0"%"'
        
        if rel_acc_val >= 90.0:
            cell_rel.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        elif rel_acc_val >= 75.0:
            cell_rel.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        else:
            cell_rel.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
            
        # RIGHT & WRONG Counts
        cell_right = ws1.cell(row=row_num, column=6, value=m["right_count"])
        cell_right.alignment = Alignment(horizontal="center", vertical="center")
        cell_right.font = font_green
        
        cell_wrong = ws1.cell(row=row_num, column=7, value=m["wrong_count"])
        cell_wrong.alignment = Alignment(horizontal="center", vertical="center")
        cell_wrong.font = font_red if m["wrong_count"] > 0 else font_body
        
        cell_w_details = ws1.cell(row=row_num, column=8, value=m["wrong_summary"])
        cell_w_details.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell_w_details.font = font_red if m["wrong_count"] > 0 else font_body
        
        ws1.cell(row=row_num, column=9, value=m["furn_aabb_overlaps"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=10, value=m["total_obb_overlap_area"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=11, value=m["oob_count"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=12, value=m["max_oob_dist"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=13, value=m["free_floor_pct"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=14, value=m["walkable_components"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=15, value=m["center_clearance"]).alignment = Alignment(horizontal="center", vertical="center")
        
        cell_status = ws1.cell(row=row_num, column=16, value=m["status"])
        cell_status.alignment = Alignment(horizontal="center", vertical="center")
        cell_status.font = font_bold
        
        if m["status"] == "EXCELLENT":
            cell_status.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="006100")
        elif m["status"] == "GOOD":
            cell_status.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="375623")
        elif m["status"] == "MINOR ISSUES":
            cell_status.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="9C6500")
        else:
            cell_status.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="9C0006")
            
        for c in range(1, 17):
            cell_item = ws1.cell(row=row_num, column=c)
            cell_item.border = border_thin
            if c not in [4, 5, 16]:
                cell_item.font = font_body
            if idx % 2 == 0 and c not in [4, 5, 16]:
                cell_item.fill = zebra_fill

    # ==========================================
    # SHEET 2: Detailed Relational Breakdown
    # ==========================================
    ws2 = wb.create_sheet(title="Relational Details")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:G2")
    cell_t2 = ws2["A1"]
    cell_t2.value = "  Detailed Spatial Relationship Verification Breakdown (RIGHT vs WRONG)"
    cell_t2.font = font_title
    cell_t2.fill = header_fill
    cell_t2.alignment = Alignment(vertical="center", horizontal="left")
    
    headers2 = ["Rank", "Log File Name", "Source Object (s)", "Relation Type", "Target Object (o)", "Status", "Measurement Details"]
    ws2.row_dimensions[4].height = 24
    for c_idx, h_text in enumerate(headers2, 1):
        c = ws2.cell(row=4, column=c_idx, value=h_text)
        c.font = font_header
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_thin
        
    col_widths2 = {"A": 8, "B": 45, "C": 22, "D": 20, "E": 22, "F": 14, "G": 30}
    for col_letter, width in col_widths2.items():
        ws2.column_dimensions[col_letter].width = width

    curr_row2 = 5
    for idx, item in enumerate(ranked_data, 1):
        fname = os.path.basename(item["file_path"])
        rel_details = item["metrics"]["rel_details"]
        
        if not rel_details:
            c_rank = ws2.cell(row=curr_row2, column=1, value=f"#{idx}")
            c_file = ws2.cell(row=curr_row2, column=2, value=fname)
            c_note = ws2.cell(row=curr_row2, column=3, value="No explicit GT relationships evaluated")
            for c in range(1, 8):
                ws2.cell(row=curr_row2, column=c).border = border_thin
            curr_row2 += 1
            continue
            
        for r_item in rel_details:
            ws2.cell(row=curr_row2, column=1, value=f"#{idx}").alignment = Alignment(horizontal="center", vertical="center")
            ws2.cell(row=curr_row2, column=2, value=fname).alignment = Alignment(horizontal="left", vertical="center")
            ws2.cell(row=curr_row2, column=3, value=f"{r_item['src_name']} ({r_item['src_id']})").alignment = Alignment(horizontal="center", vertical="center")
            ws2.cell(row=curr_row2, column=4, value=r_item['relation']).alignment = Alignment(horizontal="center", vertical="center")
            ws2.cell(row=curr_row2, column=5, value=f"{r_item['dest_name']} ({r_item['dest_id']})").alignment = Alignment(horizontal="center", vertical="center")
            
            c_status = ws2.cell(row=curr_row2, column=6, value="RIGHT" if r_item['passed'] else "WRONG")
            c_status.alignment = Alignment(horizontal="center", vertical="center")
            if r_item['passed']:
                c_status.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                c_status.font = font_green
            else:
                c_status.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                c_status.font = font_red
                
            ws2.cell(row=curr_row2, column=7, value=r_item['detail']).alignment = Alignment(horizontal="center", vertical="center")
            
            for c in range(1, 8):
                ws2.cell(row=curr_row2, column=c).border = border_thin
                if c not in [6]:
                    ws2.cell(row=curr_row2, column=c).font = font_body
            curr_row2 += 1

    wb.save(output_excel)
    print(f"Excel report successfully generated: {output_excel}")

def main():
    debug_dir = os.path.dirname(os.path.abspath(__file__))
    render_dir = os.path.join(debug_dir, "rendered_plots")
    excel_file = os.path.join(debug_dir, "layout_ranking_report.xlsx")
    target_scene = detect_target_scene(debug_dir)
    
    files = sorted(glob.glob(os.path.join(debug_dir, "*.txt")), key=natural_sort_key)
    print(f"Analyzing {len(files)} log files for scene [{target_scene}]...")
    
    scenarios = []
    for fpath in files:
        objs, overlaps, log_acc, log_mom = parse_log_file(fpath)
        if not objs:
            continue
            
        fname = os.path.basename(fpath)
        rel_details = check_file_relations(objs, target_scene)
        
        if "debug_bbox_input" in fname:
            rel_acc_val = 100.0
        elif log_acc is not None:
            rel_acc_val = log_acc
        else:
            right_c = sum(1 for r in rel_details if r['passed'])
            tot_c = len(rel_details)
            rel_acc_val = (right_c / tot_c * 100.0) if tot_c > 0 else 100.0
            
        metrics = analyze_scene(objs, overlaps, rel_acc_val, rel_details)
        
        img_fname = fname.replace(".txt", "")
        img_out = os.path.join(render_dir, f"{img_fname}.png")
        render_scene_image(fpath, objs, metrics, img_out)
        
        scenarios.append({
            "file_path": fpath,
            "metrics": metrics,
            "img_path": img_out
        })
        
    scenarios.sort(key=lambda x: (x["metrics"]["score"], x["metrics"]["relational_acc"]), reverse=True)
    
    print("\n" + "="*85)
    print(" TOP BEST LAYOUT CONFIGURATIONS ")
    print("="*85)
    for rank, item in enumerate(scenarios[:5], 1):
        m = item["metrics"]
        print(f"Rank #{rank}: {os.path.basename(item['file_path'])}")
        print(f"   Score: {m['score']:.1f}/100 [{m['status']}] | Rel Acc: {m['relational_acc']:.1f}% | RIGHT: {m['right_count']} | WRONG: {m['wrong_count']} | Overlaps: {m['furn_aabb_overlaps']} | Out-of-Bounds: {m['oob_count']}")
        if m['wrong_count'] > 0:
            print(f"   ❌ WRONG Relations: {m['wrong_summary']}")
    print("="*85)
    
    build_excel_report(scenarios, excel_file)

if __name__ == "__main__":
    main()
