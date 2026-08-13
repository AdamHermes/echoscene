import os
import glob
import re
import json
import math
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from shapely.geometry import Polygon, MultiPolygon, Point
from shapely.ops import unary_union
from scipy.spatial import ConvexHull
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenPyxlImage

BASE_DIR = "/Users/lehoangan/Documents/GitHub/ROOM/echoscene"
REL_DATA_ALL = {}
for rf in ['FRONT/relationships_bedroom_test.json', 'FRONT/relationships_diningroom_test.json', 'FRONT/relationships_all_test.json']:
    p = os.path.join(BASE_DIR, rf)
    if os.path.exists(p):
        with open(p, 'r') as f:
            data = json.load(f)
            for scan in data.get('scans', []):
                REL_DATA_ALL[scan['scan']] = scan

CLASS_NAME_MAP = {
    1: 'bed', 2: 'nightstand', 3: 'wardrobe', 4: 'chair', 5: 'table',
    6: 'chair', 7: 'lamp', 8: 'nightstand', 11: 'table', 12: 'tv_stand',
    13: 'wardrobe', 14: 'floor'
}

def natural_sort_key(s):
    """Sorts strings using natural numerical ordering (matching VS Code order)."""
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', s)]

def detect_target_scene(debug_dir):
    """Auto-detects the target scene ID from txt filenames or log content."""
    for fpath in glob.glob(os.path.join(debug_dir, "*.txt")):
        fname = os.path.basename(fpath)
        m = re.search(r'debug_bbox_input_([A-Za-z0-9\-]+)\.txt', fname)
        if m:
            return m.group(1)
        with open(fpath, 'r') as f:
            for line in f:
                m2 = re.search(r'SCENE:\s*([A-Za-z0-9\-]+)', line)
                if m2:
                    return m2.group(1)
    return "MasterBedroom-17206"

def get_obb_polygon(x, z, l, w, angle_deg):
    """Calculates 2D Shapely Polygon for Oriented Bounding Box."""
    angle_rad = np.radians(angle_deg)
    cos_a, sin_a = np.cos(angle_rad), np.sin(angle_rad)
    corners = np.array([[-l/2, -w/2], [l/2, -w/2], [l/2, w/2], [-l/2, w/2]])
    rotated = []
    for cx, cz in corners:
        rx = cx * cos_a - cz * sin_a
        rz = cx * sin_a + cz * cos_a
        rotated.append((x + rx, z + rz))
    return Polygon(rotated)

# =========================================================================
# OFFICIAL 3D-FRONT RELATIONAL EVALUATION FUNCTIONS (from helpers/metrics_3dfront.py)
# =========================================================================
def close_dis(corners1, corners2):
    dist = -2 * np.matmul(corners1, corners2.transpose())
    dist += np.sum(corners1 ** 2, axis=-1)[:, None]
    dist += np.sum(corners2 ** 2, axis=-1)[None, :]
    dist = np.sqrt(dist)
    return np.min(dist)

def cal_l2_distance(point_1, point_2):
    return np.sqrt((point_2[0] - point_1[0])**2 + (point_2[1] - point_1[1])**2)

def poly_area(x, y):
    return 0.5 * np.abs(np.dot(x, np.roll(y, 1)) - np.dot(y, np.roll(x, 1)))

def box3d_vol(corners):
    a = np.sqrt(np.sum((corners[0,:] - corners[1,:])**2))
    b = np.sqrt(np.sum((corners[1,:] - corners[2,:])**2))
    c = np.sqrt(np.sum((corners[0,:] - corners[4,:])**2))
    return a*b*c

def polygon_clip(subjectPolygon, clipPolygon):
    def inside(p):
        return (cp2[0]-cp1[0])*(p[1]-cp1[1]) > (cp2[1]-cp1[1])*(p[0]-cp1[0])
    def computeIntersection():
        dc = [ cp1[0] - cp2[0], cp1[1] - cp2[1] ]
        dp = [ s[0] - e[0], s[1] - e[1] ]
        n1 = cp1[0] * cp2[1] - cp1[1] * cp2[0]
        n2 = s[0] * e[1] - s[1] * e[0]
        n3 = 1.0 / (dc[0] * dp[1] - dc[1] * dp[0])
        return [(n1*dp[0] - n2*dc[0]) * n3, (n1*dp[1] - n2*dc[1]) * n3]
    outputList = subjectPolygon
    cp1 = clipPolygon[-1]
    for clipVertex in clipPolygon:
        cp2 = clipVertex
        inputList = outputList
        outputList = []
        s = inputList[-1]
        for subjectVertex in inputList:
            e = subjectVertex
            if inside(e):
                if not inside(s): outputList.append(computeIntersection())
                outputList.append(e)
            elif inside(s):
                outputList.append(computeIntersection())
            s = e
        cp1 = cp2
        if len(outputList) == 0: return None
    return outputList

def convex_hull_intersection(p1, p2):
    inter_p = polygon_clip(p1, p2)
    if inter_p is not None:
        hull_inter = ConvexHull(inter_p)
        return inter_p, hull_inter.volume
    else:
        return None, 0.0

def corners_from_box(box, param6=True, with_translation=False):
    if param6: l, h, w, px, py, pz = box
    else: l, h, w, px, py, pz, _ = box
    (tx, ty, tz) = (px, py, pz) if with_translation else (0,0,0)
    x_corners = [w/2,w/2,-w/2,-w/2,w/2,w/2,-w/2,-w/2]
    y_corners = [h,h,h,h,0,0,0,0]
    z_corners = [l/2,-l/2,-l/2,l/2,l/2,-l/2,-l/2,l/2]
    corners_3d = np.dot(np.eye(3), np.vstack([x_corners,y_corners,z_corners]))
    corners_3d[0,:] += tx
    corners_3d[1,:] += ty
    corners_3d[2,:] += tz
    return np.transpose(corners_3d)

def box3d_iou(box1, box2, param6=True, with_translation=False):
    corners1 = corners_from_box(box1, param6, with_translation)
    corners2 = corners_from_box(box2, param6, with_translation)
    rect1 = [(corners1[i,2], corners1[i,0]) for i in range(0,4)]
    rect2 = [(corners2[i,2], corners2[i,0]) for i in range(0,4)]
    area1 = poly_area(np.array(rect1)[:,0], np.array(rect1)[:,1])
    area2 = poly_area(np.array(rect2)[:,0], np.array(rect2)[:,1])
    inter, inter_area = convex_hull_intersection(rect1, rect2)
    iou_2d = inter_area/(area1+area2-inter_area)
    ymax = min(corners1[0,1], corners2[0,1])
    ymin = max(corners1[4,1], corners2[4,1])
    inter_vol = inter_area * max(0.0, ymax-ymin)
    vol1 = box3d_vol(corners1)
    vol2 = box3d_vol(corners2)
    volmin = min(vol1, vol2)
    iou = inter_vol / volmin
    return iou, iou_2d

def check_file_relations(objects, target_scene, strict=True, overlap_threshold=0.3):
    """Evaluates individual GT spatial relationships using official 3D-FRONT criteria."""
    if not target_scene or target_scene not in REL_DATA_ALL or not objects:
        return []
        
    scan_rel = REL_DATA_ALL[target_scene]
    rel_objects = scan_rel['objects']
    relationships = scan_rel['relationships']
    
    furn_objs = [o for o in objects if o['name'] not in ['_scene_']]
    pred_boxes = np.array([[o['l'], o['h'], o['w'], o['x'], o['y'], o['z']] for o in furn_objs])
    param6 = True
    
    keys_sorted = list(rel_objects.keys())
    results = []
    
    for rel in relationships:
        src, dest, rel_idx, rel_type = rel
        src_str, dest_str = str(src), str(dest)
        
        if src_str not in keys_sorted or dest_str not in keys_sorted:
            continue
            
        s_i = keys_sorted.index(src_str)
        d_i = keys_sorted.index(dest_str)
        
        if s_i >= len(furn_objs) or d_i >= len(furn_objs):
            continue
            
        obj_s = furn_objs[s_i]
        obj_d = furn_objs[d_i]
        
        if obj_s['name'] == 'floor' or obj_d['name'] == 'floor':
            continue
            
        box_s = pred_boxes[s_i]
        box_o = pred_boxes[d_i]
        
        passed = False
        val_str = ''
        
        if rel_type == 'left':
            diff = box_s[5] - box_o[5]
            iou_val = box3d_iou(box_s, box_o, param6=param6, with_translation=True)[0]
            passed = not (diff > -0.05 or (strict and iou_val > overlap_threshold))
            val_str = f'dz={diff:+.3f}m, iou={iou_val:.2f}'
        elif rel_type == 'right':
            diff = box_s[5] - box_o[5]
            iou_val = box3d_iou(box_s, box_o, param6=param6, with_translation=True)[0]
            passed = not (diff < 0.05 or (strict and iou_val > overlap_threshold))
            val_str = f'dz={diff:+.3f}m, iou={iou_val:.2f}'
        elif rel_type == 'front':
            diff = box_s[3] - box_o[3]
            iou_val = box3d_iou(box_s, box_o, param6=param6, with_translation=True)[0]
            passed = not (diff < -0.05 or (strict and iou_val > overlap_threshold))
            val_str = f'dx={diff:+.3f}m, iou={iou_val:.2f}'
        elif rel_type == 'behind':
            diff = box_s[3] - box_o[3]
            iou_val = box3d_iou(box_s, box_o, param6=param6, with_translation=True)[0]
            passed = not (diff > 0.05 or (strict and iou_val > overlap_threshold))
            val_str = f'dx={diff:+.3f}m, iou={iou_val:.2f}'
        elif rel_type == 'bigger than':
            vol_s = box_s[0] * box_s[1] * box_s[2]
            vol_o = box_o[0] * box_o[1] * box_o[2]
            ratio = (vol_s - vol_o) / vol_s
            passed = not (ratio < 0.15)
            val_str = f'ratio={ratio:+.2f} (vol_s={vol_s:.2f}m³ vs vol_o={vol_o:.2f}m³)'
        elif rel_type == 'smaller than':
            vol_s = box_s[0] * box_s[1] * box_s[2]
            vol_o = box_o[0] * box_o[1] * box_o[2]
            ratio = (vol_s - vol_o) / vol_s
            passed = not (ratio > -0.15)
            val_str = f'ratio={ratio:+.2f} (vol_s={vol_s:.2f}m³ vs vol_o={vol_o:.2f}m³)'
        elif rel_type == 'taller than':
            absheight_s = box_s[4] + box_s[1]
            absheight_o = box_o[4] + box_o[1]
            ratio = (absheight_s - absheight_o) / absheight_s
            passed = not (ratio < 0.1)
            val_str = f'height_diff={absheight_s-absheight_o:+.2f}m'
        elif rel_type == 'shorter than':
            absheight_s = box_s[4] + box_s[1]
            absheight_o = box_o[4] + box_o[1]
            ratio = (absheight_s - absheight_o) / absheight_s
            passed = not (ratio > -0.1)
            val_str = f'height_diff={absheight_s-absheight_o:+.2f}m'
        elif rel_type in ['above', 'standing on']:
            diff = abs(box_s[4] - box_o[4])
            passed = diff < 0.04
            val_str = f'dy={diff:.3f}m'
        elif rel_type == 'close by':
            corners_s = corners_from_box(box_s, param6, with_translation=True)
            corners_o = corners_from_box(box_o, param6, with_translation=True)
            c_dist1 = close_dis(corners_s, corners_o)
            passed = c_dist1 <= 0.45
            val_str = f'3d_dist={c_dist1:.3f}m'
        elif rel_type == 'symmetrical to':
            s_flip_xz = [-box_s[3], -box_s[5]]
            s_flip_x = [-box_s[3], box_s[5]]
            s_flip_z = [box_s[3], -box_s[5]]
            o_center = [box_o[3], box_o[5]]
            d1 = cal_l2_distance(s_flip_xz, o_center)
            d2 = cal_l2_distance(s_flip_x, o_center)
            d3 = cal_l2_distance(s_flip_z, o_center)
            min_d = min(d1, d2, d3)
            passed = min_d < 0.45
            val_str = f'symm_dist={min_d:.3f}m'
        else:
            passed = True
            val_str = 'ok'
            
        results.append({
            'src_name': obj_s['name'],
            'src_id': src_str,
            'dest_name': obj_d['name'],
            'dest_id': dest_str,
            'relation': rel_type,
            'passed': passed,
            'detail': val_str
        })
    return results

def parse_log_file(file_path):
    """Parses raw object bounding box details, pairwise overlaps, and explicit relational metrics from a log file."""
    objects = []
    raw_overlaps = []
    rel_acc_log = None
    means_of_mean_log = None
    
    if not os.path.exists(file_path):
        return objects, raw_overlaps, rel_acc_log, means_of_mean_log
        
    with open(file_path, 'r') as f:
        lines = f.readlines()
        
    content = "".join(lines)
    
    acc_m = re.search(r'acc &.*Total:\s*&?\s*([\d\.]+)', content)
    if acc_m:
        val = float(acc_m.group(1))
        rel_acc_log = val * 100.0 if val <= 1.0 else val
        
    means_m = re.search(r'means of mean:\s*([\d\.]+)', content)
    if means_m:
        val = float(means_m.group(1))
        means_of_mean_log = val * 100.0 if val <= 1.0 else val

    start_parsing = False
    for line in lines:
        if "---" in line:
            start_parsing = True
            continue
        if "Pairwise overlap" in line or "==" in line:
            start_parsing = False
        if "!! OVERLAP:" in line:
            raw_overlaps.append(line.strip())
            
        if start_parsing and line.strip():
            parts = line.split()
            if len(parts) >= 8:
                name = parts[0]
                try:
                    l = float(parts[1])
                    h = float(parts[2])
                    w = float(parts[3])
                    x = float(parts[4])
                    y = float(parts[5])
                    z = float(parts[6])
                    angle = float(parts[7].replace('°', ''))
                    objects.append({
                        "name": name, "l": l, "h": h, "w": w,
                        "x": x, "y": y, "z": z, "angle": angle
                    })
                except ValueError:
                    continue
    return objects, raw_overlaps, rel_acc_log, means_of_mean_log

def parse_json_scene_objects(json_path, scene_id):
    """Extracts object list from JSON models (Baseline / Released Full Model)."""
    if not os.path.exists(json_path):
        return []
    with open(json_path, 'r') as f:
        data = json.load(f)
    if scene_id not in data['scene_ids']:
        return []
    idx = data['scene_ids'].index(scene_id)
    
    class_labels = np.array(data['class_labels'][idx])
    translations = np.array(data['translations'][idx])
    sizes = np.array(data['sizes'][idx])
    angles = np.array(data['angles'][idx])
    objectness = np.array(data['objectness'][idx]).squeeze()
    
    active_mask = objectness > 0.5
    active_indices = np.where(active_mask)[0]
    
    objects = []
    for i, a_idx in enumerate(active_indices):
        c_idx = int(np.argmax(class_labels[a_idx]))
        c_name = CLASS_NAME_MAP.get(c_idx, 'furniture')
        l, h, w = sizes[a_idx]
        x, y, z = translations[a_idx]
        angle_deg = float(np.degrees(angles[a_idx][0]))
        objects.append({
            'name': c_name, 'l': float(l), 'h': float(h), 'w': float(w),
            'x': float(x), 'y': float(y), 'z': float(z), 'angle': angle_deg
        })
    if not any(o['name'] == 'floor' for o in objects):
        objects.append({'name': 'floor', 'l': 4.0, 'h': 0.0, 'w': 3.5, 'x': 0.0, 'y': 0.0, 'z': 0.0, 'angle': 0.0})
    return objects

def analyze_scene(objects, raw_overlaps, rel_acc_val, rel_details):
    """Computes layout quality metrics (collisions, OBB overlap area, out-of-bounds, walkability, relational acc)."""
    floor_obj = None
    furniture_objs = []
    for obj in objects:
        if obj["name"] == "floor":
            floor_obj = obj
        elif obj["name"] not in ["_scene_", "lamp"]:
            furniture_objs.append(obj)
            
    if floor_obj is not None:
        floor_poly = get_obb_polygon(floor_obj["x"], floor_obj["z"], floor_obj["l"], floor_obj["w"], floor_obj["angle"])
    else:
        floor_poly = Polygon([[-2.0, -2.5], [2.0, -2.5], [2.0, 2.5], [-2.0, 2.5]])
        
    furn_polys = []
    for obj in furniture_objs:
        poly = get_obb_polygon(obj["x"], obj["z"], obj["l"], obj["w"], obj["angle"])
        furn_polys.append((obj, poly))
        
    furn_aabb_overlaps = 0
    for line in raw_overlaps:
        match = re.search(r'OVERLAP:\s*(\w+)\s*<->\s*(\w+)', line)
        if match:
            o1, o2 = match.group(1), match.group(2)
            if o1 not in ["floor", "lamp", "_scene_"] and o2 not in ["floor", "lamp", "_scene_"]:
                furn_aabb_overlaps += 1
                
    obb_overlap_count = 0
    total_obb_overlap_area = 0.0
    N = len(furn_polys)
    for i in range(N):
        for j in range(i+1, N):
            p1 = furn_polys[i][1]
            p2 = furn_polys[j][1]
            if p1.intersects(p2):
                inter = p1.intersection(p2)
                if inter.area > 1e-4:
                    obb_overlap_count += 1
                    total_obb_overlap_area += inter.area
                    
    oob_count = 0
    max_oob_dist = 0.0
    total_oob_area = 0.0
    
    for obj, poly in furn_polys:
        if not floor_poly.contains(poly):
            diff = poly.difference(floor_poly)
            if diff.area > 1e-4:
                oob_count += 1
                total_oob_area += diff.area
                coords = np.array(poly.exterior.coords)
                for cx, cz in coords:
                    pt_dist = floor_poly.distance(Point(cx, cz))
                    if pt_dist > max_oob_dist:
                        max_oob_dist = pt_dist

    if furn_polys:
        all_furn_union = unary_union([p for _, p in furn_polys])
        free_floor = floor_poly.difference(all_furn_union)
    else:
        free_floor = floor_poly
        
    free_floor_area = free_floor.area
    free_floor_pct = (free_floor_area / max(floor_poly.area, 1e-5)) * 100.0
    
    if isinstance(free_floor, Polygon):
        walkable_components = 1 if free_floor.area > 0.01 else 0
    elif isinstance(free_floor, MultiPolygon):
        walkable_components = sum(1 for p in free_floor.geoms if p.area > 0.05)
    else:
        walkable_components = 0
        
    center_pt = Point(0, 0)
    min_center_dist = 999.0
    for _, poly in furn_polys:
        d = poly.distance(center_pt)
        if d < min_center_dist:
            min_center_dist = d
    if min_center_dist == 999.0:
        min_center_dist = 0.0

    right_count = sum(1 for r in rel_details if r['passed'])
    wrong_count = sum(1 for r in rel_details if not r['passed'])
    wrong_summary = ", ".join([f"{r['src_name']}({r['src_id']})--[{r['relation']}]-->{r['dest_name']}({r['dest_id']})" for r in rel_details if not r['passed']])
    if not wrong_summary:
        wrong_summary = "None (All Relations RIGHT)"

    score = 100.0
    score -= furn_aabb_overlaps * 20.0
    score -= total_obb_overlap_area * 35.0
    score -= oob_count * 15.0
    score -= max_oob_dist * 25.0
    score -= max(0, walkable_components - 1) * 8.0
    score += min(10.0, free_floor_pct * 0.15)
    score = float(np.clip(score, 0.0, 100.0))
    
    if score >= 90.0:
        status = "EXCELLENT"
    elif score >= 75.0:
        status = "GOOD"
    elif score >= 50.0:
        status = "MINOR ISSUES"
    else:
        status = "SEVERE COLLISIONS"
        
    return {
        "score": score,
        "status": status,
        "relational_acc": round(rel_acc_val, 1),
        "right_count": right_count,
        "wrong_count": wrong_count,
        "wrong_summary": wrong_summary,
        "rel_details": rel_details,
        "furn_aabb_overlaps": furn_aabb_overlaps,
        "obb_overlap_count": obb_overlap_count,
        "total_obb_overlap_area": round(total_obb_overlap_area, 4),
        "oob_count": oob_count,
        "max_oob_dist": round(max_oob_dist, 4),
        "free_floor_pct": round(free_floor_pct, 1),
        "walkable_components": walkable_components,
        "center_clearance": round(min_center_dist, 3)
    }

def render_scene_image(file_label, objects, metrics, out_img_path):
    """Renders high-quality 2D top-down layout plot image with metric overlay."""
    fig, ax = plt.subplots(figsize=(5.5, 5.5), dpi=150)
    
    color_map = {
        "bed": "#3498db", "chair": "#e74c3c", "nightstand": "#9b59b6", 
        "table": "#e67e22", "wardrobe": "#2ecc71", "lamp": "#f1c40f", 
        "floor": "#ecf0f1"
    }
    
    floor_obj = next((o for o in objects if o["name"] == "floor"), None)
    if floor_obj:
        f_poly = get_obb_polygon(floor_obj["x"], floor_obj["z"], floor_obj["l"], floor_obj["w"], floor_obj["angle"])
        f_patch = patches.Polygon(np.array(f_poly.exterior.coords), closed=True, facecolor="#f8f9fa", edgecolor="#7f8c8d", linewidth=2.0)
        ax.add_patch(f_patch)
        
    for obj in objects:
        if obj["name"] in ["_scene_", "floor"]:
            continue
        poly = get_obb_polygon(obj["x"], obj["z"], obj["l"], obj["w"], obj["angle"])
        c = color_map.get(obj["name"], "#1abc9c")
        alpha = 0.7 if obj["name"] != "lamp" else 0.4
        
        patch = patches.Polygon(np.array(poly.exterior.coords), closed=True, facecolor=c, edgecolor="#2c3e50", alpha=alpha, linewidth=1.5)
        ax.add_patch(patch)
        
        if obj["name"] != "lamp":
            ax.text(obj["x"], obj["z"], obj["name"], ha='center', va='center', fontsize=8, weight='bold',
                    bbox=dict(facecolor='white', alpha=0.75, edgecolor='none', pad=1.5))
                    
    ax.set_aspect('equal')
    ax.set_xlim(-3.2, 3.2)
    ax.set_ylim(-3.2, 3.2)
    ax.grid(True, linestyle=':', alpha=0.4)
    
    score = metrics["score"]
    rel_acc = metrics["relational_acc"]
    status = metrics["status"]
    
    clean_label = file_label.replace("[BASELINE MODEL] ", "").replace("[RELEASED FULL MODEL] ", "")
    if len(clean_label) > 40:
        clean_label = clean_label[:37] + "..."
    ax.set_title(f"{clean_label}\nScore: {score:.1f}/100 [{status}] | Rel Acc: {rel_acc:.1f}%", fontsize=8.5, weight='bold', color='black', pad=6)
    
    plt.tight_layout()
    os.makedirs(os.path.dirname(out_img_path), exist_ok=True)
    plt.savefig(out_img_path, dpi=150, bbox_inches='tight')
    plt.close(fig)

def build_excel_report(ranked_data, output_excel):
    """Creates a formatted Excel report with embedded layout images, summary table, and detailed relation breakdown sheet."""
    wb = openpyxl.Workbook()
    
    # ==========================================
    # SHEET 1: Layout Rankings Summary
    # ==========================================
    ws1 = wb.active
    ws1.title = "Layout Rankings Summary"
    ws1.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    sub_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    special_row_fill = PatternFill(start_color="E6EEF8", end_color="E6EEF8", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
    
    font_title = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_special = Font(name="Segoe UI", size=10.5, bold=True, color="1F4E78")
    font_red = Font(name="Segoe UI", size=10, bold=True, color="9C0006")
    font_green = Font(name="Segoe UI", size=10, bold=True, color="006100")
    
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    ws1.merge_cells("A1:P2")
    cell_title = ws1["A1"]
    cell_title.value = "  ROOM Layout Quality & Relational Accuracy Benchmark Report"
    cell_title.font = font_title
    cell_title.fill = header_fill
    cell_title.alignment = Alignment(vertical="center", horizontal="left")
    
    ws1.merge_cells("A3:P3")
    cell_sub = ws1["A3"]
    cell_sub.value = f"  Total Runs Analyzed: {len(ranked_data)}  |  Includes 2 Special Benchmark Rows (Baseline Model & Complete Released Full Model)"
    cell_sub.font = Font(name="Segoe UI", size=10, italic=True, color="1F4E78")
    cell_sub.fill = sub_header_fill
    cell_sub.alignment = Alignment(vertical="center", horizontal="left")
    
    headers1 = [
        "Rank", "Layout Visualization", "Log File / Model Name", "Quality Score", "Relational Acc (%)",
        "RIGHT Relations", "WRONG Relations", "WRONG Relation Details",
        "Overlaps (Count)", "OBB Overlap (m²)", "Out of Bounds", "Max OOB Dist (m)",
        "Free Floor %", "Walkable Components", "Center Clear (m)", "Grade Status"
    ]
    
    ws1.row_dimensions[5].height = 28
    for col_num, h_text in enumerate(headers1, 1):
        cell = ws1.cell(row=5, column=col_num)
        cell.value = h_text
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    col_widths1 = {
        "A": 18,  "B": 24,  "C": 48,  "D": 15,  "E": 18,  "F": 16,
        "G": 16,  "H": 40,  "I": 16,  "J": 18,  "K": 15,  "L": 18,
        "M": 14,  "N": 20,  "O": 16,  "P": 20
    }
    for col_letter, width in col_widths1.items():
        ws1.column_dimensions[col_letter].width = width

    start_row = 6
    for idx, item in enumerate(ranked_data, 1):
        row_num = start_row + idx - 1
        ws1.row_dimensions[row_num].height = 115
        
        m = item["metrics"]
        fname = item["display_name"]
        img_path = item["img_path"]
        is_special = item.get("is_special", False)
        
        rank_str = item.get("rank_label", f"#{idx}")
        cell_rank = ws1.cell(row=row_num, column=1, value=rank_str)
        cell_rank.alignment = Alignment(horizontal="center", vertical="center")
        cell_rank.font = font_special if is_special else font_bold
        
        if os.path.exists(img_path):
            img = OpenPyxlImage(img_path)
            img.width = 140
            img.height = 140
            ws1.add_image(img, f"B{row_num}")
            
        cell_file = ws1.cell(row=row_num, column=3, value=fname)
        cell_file.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell_file.font = font_special if is_special else font_body
        
        score_val = m["score"]
        cell_score = ws1.cell(row=row_num, column=4, value=score_val)
        cell_score.alignment = Alignment(horizontal="center", vertical="center")
        cell_score.font = Font(name="Segoe UI", size=11, bold=True)
        cell_score.number_format = '0.0'
        
        if score_val >= 90:
            cell_score.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        elif score_val >= 75:
            cell_score.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        elif score_val >= 50:
            cell_score.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        else:
            cell_score.fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
            
        rel_acc_fraction = m["relational_acc"] / 100.0
        cell_rel = ws1.cell(row=row_num, column=5, value=rel_acc_fraction)
        cell_rel.alignment = Alignment(horizontal="center", vertical="center")
        cell_rel.font = Font(name="Segoe UI", size=11, bold=True)
        cell_rel.number_format = '0.0%'
        
        if m["relational_acc"] >= 90.0:
            cell_rel.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        elif m["relational_acc"] >= 75.0:
            cell_rel.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        else:
            cell_rel.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
            
        cell_right = ws1.cell(row=row_num, column=6, value=m["right_count"])
        cell_right.alignment = Alignment(horizontal="center", vertical="center")
        cell_right.font = font_green
        
        cell_wrong = ws1.cell(row=row_num, column=7, value=m["wrong_count"])
        cell_wrong.alignment = Alignment(horizontal="center", vertical="center")
        cell_wrong.font = font_red if m["wrong_count"] > 0 else font_body
        
        cell_w_details = ws1.cell(row=row_num, column=8, value=m["wrong_summary"])
        cell_w_details.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell_w_details.font = font_red if m["wrong_count"] > 0 else font_body
        
        ws1.cell(row=row_num, column=9, value=m["furn_aabb_overlaps"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=10, value=m["total_obb_overlap_area"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=11, value=m["oob_count"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=12, value=m["max_oob_dist"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=13, value=m["free_floor_pct"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=14, value=m["walkable_components"]).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=row_num, column=15, value=m["center_clearance"]).alignment = Alignment(horizontal="center", vertical="center")
        
        cell_status = ws1.cell(row=row_num, column=16, value=m["status"])
        cell_status.alignment = Alignment(horizontal="center", vertical="center")
        cell_status.font = font_bold
        
        if m["status"] == "EXCELLENT":
            cell_status.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="006100")
        elif m["status"] == "GOOD":
            cell_status.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="375623")
        elif m["status"] == "MINOR ISSUES":
            cell_status.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="9C6500")
        else:
            cell_status.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="9C0006")
            
        for c in range(1, 17):
            cell_item = ws1.cell(row=row_num, column=c)
            cell_item.border = border_thin
            if c not in [4, 5, 16]:
                cell_item.font = font_special if is_special else font_body
            if is_special and c not in [4, 5, 16]:
                cell_item.fill = special_row_fill
            elif idx % 2 == 0 and c not in [4, 5, 16]:
                cell_item.fill = zebra_fill

    # ==========================================
    # SHEET 2: Detailed Relational Breakdown
    # ==========================================
    ws2 = wb.create_sheet(title="Relational Details")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:G2")
    cell_t2 = ws2["A1"]
    cell_t2.value = "  Detailed Spatial Relationship Verification Breakdown (RIGHT vs WRONG)"
    cell_t2.font = font_title
    cell_t2.fill = header_fill
    cell_t2.alignment = Alignment(vertical="center", horizontal="left")
    
    headers2 = ["Rank", "Log File / Model Name", "Source Object (s)", "Relation Type", "Target Object (o)", "Status", "Measurement Details"]
    ws2.row_dimensions[4].height = 24
    for c_idx, h_text in enumerate(headers2, 1):
        c = ws2.cell(row=4, column=c_idx, value=h_text)
        c.font = font_header
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_thin
        
    col_widths2 = {"A": 18, "B": 48, "C": 22, "D": 20, "E": 22, "F": 14, "G": 30}
    for col_letter, width in col_widths2.items():
        ws2.column_dimensions[col_letter].width = width

    curr_row2 = 5
    for idx, item in enumerate(ranked_data, 1):
        fname = item["display_name"]
        rel_details = item["metrics"]["rel_details"]
        rank_str = item.get("rank_label", f"#{idx}")
        is_special = item.get("is_special", False)
        
        if not rel_details:
            ws2.cell(row=curr_row2, column=1, value=rank_str)
            ws2.cell(row=curr_row2, column=2, value=fname)
            ws2.cell(row=curr_row2, column=3, value="No explicit GT relationships evaluated")
            for c in range(1, 8):
                ws2.cell(row=curr_row2, column=c).border = border_thin
            curr_row2 += 1
            continue
            
        for r_item in rel_details:
            ws2.cell(row=curr_row2, column=1, value=rank_str).alignment = Alignment(horizontal="center", vertical="center")
            ws2.cell(row=curr_row2, column=2, value=fname).alignment = Alignment(horizontal="left", vertical="center")
            ws2.cell(row=curr_row2, column=3, value=f"{r_item['src_name']} ({r_item['src_id']})").alignment = Alignment(horizontal="center", vertical="center")
            ws2.cell(row=curr_row2, column=4, value=r_item['relation']).alignment = Alignment(horizontal="center", vertical="center")
            ws2.cell(row=curr_row2, column=5, value=f"{r_item['dest_name']} ({r_item['dest_id']})").alignment = Alignment(horizontal="center", vertical="center")
            
            c_status = ws2.cell(row=curr_row2, column=6, value="RIGHT" if r_item['passed'] else "WRONG")
            c_status.alignment = Alignment(horizontal="center", vertical="center")
            if r_item['passed']:
                c_status.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                c_status.font = font_green
            else:
                c_status.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                c_status.font = font_red
                
            ws2.cell(row=curr_row2, column=7, value=r_item['detail']).alignment = Alignment(horizontal="center", vertical="center")
            
            for c in range(1, 8):
                ws2.cell(row=curr_row2, column=c).border = border_thin
                if c not in [6]:
                    ws2.cell(row=curr_row2, column=c).font = font_special if is_special else font_body
                if is_special and c not in [6]:
                    ws2.cell(row=curr_row2, column=c).fill = special_row_fill
            curr_row2 += 1

    wb.save(output_excel)
    print(f"Excel report successfully generated: {output_excel}")

def main():
    debug_dir = os.path.dirname(os.path.abspath(__file__))
    render_dir = os.path.join(debug_dir, "rendered_plots")
    excel_file = os.path.join(debug_dir, "layout_ranking_report.xlsx")
    target_scene = detect_target_scene(debug_dir)
    
    files = sorted(glob.glob(os.path.join(debug_dir, "*.txt")), key=natural_sort_key)
    print(f"Analyzing {len(files)} log files for scene [{target_scene}]...")
    
    scenarios = []
    for fpath in files:
        objs, overlaps, log_acc, log_mom = parse_log_file(fpath)
        if not objs:
            continue
            
        fname = os.path.basename(fpath)
        rel_details = check_file_relations(objs, target_scene)
        
        if "debug_bbox_input" in fname:
            rel_acc_val = 100.0
        elif log_acc is not None:
            rel_acc_val = log_acc
        else:
            right_c = sum(1 for r in rel_details if r['passed'])
            tot_c = len(rel_details)
            rel_acc_val = (right_c / tot_c * 100.0) if tot_c > 0 else 100.0
            
        metrics = analyze_scene(objs, overlaps, rel_acc_val, rel_details)
        
        img_fname = fname.replace(".txt", "")
        img_out = os.path.join(render_dir, f"{img_fname}.png")
        render_scene_image(fname, objs, metrics, img_out)
        
        scenarios.append({
            "file_path": fpath,
            "display_name": fname,
            "metrics": metrics,
            "img_path": img_out,
            "is_special": False
        })
        
    scenarios.sort(key=lambda x: (x["metrics"]["score"], x["metrics"]["relational_acc"]), reverse=True)
    
    for idx, item in enumerate(scenarios, 1):
        item["rank_label"] = f"#{idx}"

    special_benchmarks = [
        ("[BASELINE MODEL] Physcene Collision Resolved", os.path.join(BASE_DIR, "baseline/physcene_collision_resolved.json"), "baseline_physcene_collision_resolved"),
        ("[RELEASED FULL MODEL] Complete Released Model Resolved", os.path.join(BASE_DIR, "current_works/to_be_merged/complete_released_full_model/vis/2050/physcene_collision_resolved.json"), "complete_released_full_model_resolved")
    ]
    
    for b_label, jpath, img_stem in special_benchmarks:
        if os.path.exists(jpath):
            objs_sp = parse_json_scene_objects(jpath, target_scene)
            if objs_sp:
                rel_sp = check_file_relations(objs_sp, target_scene)
                right_sp = sum(1 for r in rel_sp if r['passed'])
                tot_sp = len(rel_sp)
                acc_sp = (right_sp / tot_sp * 100.0) if tot_sp > 0 else 100.0
                metrics_sp = analyze_scene(objs_sp, [], acc_sp, rel_sp)
                
                img_out_sp = os.path.join(render_dir, f"{img_stem}.png")
                render_scene_image(b_label, objs_sp, metrics_sp, img_out_sp)
                
                scenarios.append({
                    "file_path": jpath,
                    "display_name": b_label,
                    "metrics": metrics_sp,
                    "img_path": img_out_sp,
                    "is_special": True,
                    "rank_label": "[BENCHMARK]"
                })
    
    print("\n" + "="*85)
    print(" ALL LAYOUT CONFIGURATIONS (INCLUDING SPECIAL BENCHMARK ROWS) ")
    print("="*85)
    for item in scenarios:
        m = item["metrics"]
        tag = item.get("rank_label", "")
        print(f"{tag:<12} {item['display_name']}")
        print(f"   Score: {m['score']:.1f}/100 [{m['status']}] | Rel Acc: {m['relational_acc']:.1f}% | RIGHT: {m['right_count']} | WRONG: {m['wrong_count']} | Overlaps: {m['furn_aabb_overlaps']} | Out-of-Bounds: {m['oob_count']}")
        if m['wrong_count'] > 0:
            print(f"   ❌ WRONG Relations: {m['wrong_summary']}")
    print("="*85)
    
    build_excel_report(scenarios, excel_file)

if __name__ == "__main__":
    main()
