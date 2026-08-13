import os
import glob
import re
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from shapely.geometry import Polygon, MultiPolygon, Point
from shapely.ops import unary_union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenPyxlImage

def natural_sort_key(s):
    """Sorts strings using natural numerical ordering (matching VS Code order)."""
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', s)]

def get_obb_polygon(x, z, l, w, angle_deg):
    """Calculates 2D Shapely Polygon for Oriented Bounding Box."""
    angle_rad = np.radians(angle_deg)
    cos_a, sin_a = np.cos(angle_rad), np.sin(angle_rad)
    corners = np.array([[-l/2, -w/2], [l/2, -w/2], [l/2, w/2], [-l/2, w/2]])
    rotated = []
    for cx, cz in corners:
        rx = cx * cos_a - cz * sin_a
        rz = cx * sin_a + cz * cos_a
        rotated.append((x + rx, z + rz))
    return Polygon(rotated)

def parse_log_file(file_path):
    """Parses raw object bounding box details and pairwise overlap logs."""
    objects = []
    raw_overlaps = []
    if not os.path.exists(file_path):
        return objects, raw_overlaps
        
    with open(file_path, 'r') as f:
        lines = f.readlines()
        
    start_parsing = False
    for line in lines:
        if "---" in line:
            start_parsing = True
            continue
        if "Pairwise overlap" in line or "==" in line:
            start_parsing = False
        if "!! OVERLAP:" in line:
            raw_overlaps.append(line.strip())
            
        if start_parsing and line.strip():
            parts = line.split()
            if len(parts) >= 8:
                name = parts[0]
                try:
                    l = float(parts[1])
                    h = float(parts[2])
                    w = float(parts[3])
                    x = float(parts[4])
                    y = float(parts[5])
                    z = float(parts[6])
                    angle = float(parts[7].replace('°', ''))
                    objects.append({
                        "name": name, "l": l, "h": h, "w": w,
                        "x": x, "y": y, "z": z, "angle": angle
                    })
                except ValueError:
                    continue
    return objects, raw_overlaps

def analyze_scene(objects, raw_overlaps):
    """Computes layout quality metrics (collisions, OBB overlap area, out-of-bounds, walkability)."""
    # Separate floor and furniture
    floor_obj = None
    furniture_objs = []
    for obj in objects:
        if obj["name"] == "floor":
            floor_obj = obj
        elif obj["name"] not in ["_scene_", "lamp"]:
            furniture_objs.append(obj)
            
    # Define Floor Polygon
    if floor_obj is not None:
        floor_poly = get_obb_polygon(floor_obj["x"], floor_obj["z"], floor_obj["l"], floor_obj["w"], floor_obj["angle"])
    else:
        floor_poly = Polygon([[-2.0, -2.5], [2.0, -2.5], [2.0, 2.5], [-2.0, 2.5]])
        
    # Furniture OBB Polygons
    furn_polys = []
    for obj in furniture_objs:
        poly = get_obb_polygon(obj["x"], obj["z"], obj["l"], obj["w"], obj["angle"])
        furn_polys.append((obj, poly))
        
    # 1. Furniture-Furniture AABB Overlaps (from log)
    furn_aabb_overlaps = 0
    for line in raw_overlaps:
        # Check if overlap does not involve floor or lamp
        match = re.search(r'OVERLAP:\s*(\w+)\s*<->\s*(\w+)', line)
        if match:
            o1, o2 = match.group(1), match.group(2)
            if o1 not in ["floor", "lamp", "_scene_"] and o2 not in ["floor", "lamp", "_scene_"]:
                furn_aabb_overlaps += 1
                
    # 2. OBB Exact 2D Overlap Area & Count
    obb_overlap_count = 0
    total_obb_overlap_area = 0.0
    N = len(furn_polys)
    for i in range(N):
        for j in range(i+1, N):
            p1 = furn_polys[i][1]
            p2 = furn_polys[j][1]
            if p1.intersects(p2):
                inter = p1.intersection(p2)
                if inter.area > 1e-4:
                    obb_overlap_count += 1
                    total_obb_overlap_area += inter.area
                    
    # 3. Room Outer Boundary Violations (Out of Bounds)
    oob_count = 0
    max_oob_dist = 0.0
    total_oob_area = 0.0
    
    for obj, poly in furn_polys:
        if not floor_poly.contains(poly):
            diff = poly.difference(floor_poly)
            if diff.area > 1e-4:
                oob_count += 1
                total_oob_area += diff.area
                # Calculate max distance corner extends outside floor
                coords = np.array(poly.exterior.coords)
                for cx, cz in coords:
                    pt_dist = floor_poly.distance(Point(cx, cz))
                    if pt_dist > max_oob_dist:
                        max_oob_dist = pt_dist

    # 4. Free Floor Area & Walkable Connectivity
    if furn_polys:
        all_furn_union = unary_union([p for _, p in furn_polys])
        free_floor = floor_poly.difference(all_furn_union)
    else:
        free_floor = floor_poly
        
    free_floor_area = free_floor.area
    free_floor_pct = (free_floor_area / max(floor_poly.area, 1e-5)) * 100.0
    
    if isinstance(free_floor, Polygon):
        walkable_components = 1 if free_floor.area > 0.01 else 0
    elif isinstance(free_floor, MultiPolygon):
        walkable_components = sum(1 for p in free_floor.geoms if p.area > 0.05)
    else:
        walkable_components = 0
        
    # 5. Center Clearance (Distance from (0,0) to nearest furniture)
    center_pt = Point(0, 0)
    min_center_dist = 999.0
    for _, poly in furn_polys:
        d = poly.distance(center_pt)
        if d < min_center_dist:
            min_center_dist = d
    if min_center_dist == 999.0:
        min_center_dist = 0.0

    # 6. Overall Quality Score (0 to 100)
    score = 100.0
    score -= furn_aabb_overlaps * 20.0
    score -= total_obb_overlap_area * 35.0
    score -= oob_count * 15.0
    score -= max_oob_dist * 25.0
    score -= max(0, walkable_components - 1) * 8.0
    score += min(10.0, free_floor_pct * 0.15)
    score = float(np.clip(score, 0.0, 100.0))
    
    # Status / Grade
    if score >= 90.0:
        status = "EXCELLENT"
    elif score >= 75.0:
        status = "GOOD"
    elif score >= 50.0:
        status = "MINOR ISSUES"
    else:
        status = "SEVERE COLLISIONS"
        
    return {
        "score": score,
        "status": status,
        "furn_aabb_overlaps": furn_aabb_overlaps,
        "obb_overlap_count": obb_overlap_count,
        "total_obb_overlap_area": round(total_obb_overlap_area, 4),
        "oob_count": oob_count,
        "max_oob_dist": round(max_oob_dist, 4),
        "free_floor_pct": round(free_floor_pct, 1),
        "walkable_components": walkable_components,
        "center_clearance": round(min_center_dist, 3)
    }

def render_scene_image(file_path, objects, metrics, out_img_path):
    """Renders high-quality 2D top-down layout plot image with metric overlay."""
    fig, ax = plt.subplots(figsize=(5.5, 5.5), dpi=150)
    
    color_map = {
        "bed": "#3498db", "chair": "#e74c3c", "nightstand": "#9b59b6", 
        "table": "#e67e22", "wardrobe": "#2ecc71", "lamp": "#f1c40f", 
        "floor": "#ecf0f1"
    }
    
    # Draw floor first
    floor_obj = next((o for o in objects if o["name"] == "floor"), None)
    if floor_obj:
        f_poly = get_obb_polygon(floor_obj["x"], floor_obj["z"], floor_obj["l"], floor_obj["w"], floor_obj["angle"])
        f_patch = patches.Polygon(np.array(f_poly.exterior.coords), closed=True, facecolor="#f8f9fa", edgecolor="#7f8c8d", linewidth=2.0)
        ax.add_patch(f_patch)
        
    # Draw furniture OBBs
    for obj in objects:
        if obj["name"] in ["_scene_", "floor"]:
            continue
        poly = get_obb_polygon(obj["x"], obj["z"], obj["l"], obj["w"], obj["angle"])
        c = color_map.get(obj["name"], "#1abc9c")
        alpha = 0.7 if obj["name"] != "lamp" else 0.4
        
        patch = patches.Polygon(np.array(poly.exterior.coords), closed=True, facecolor=c, edgecolor="#2c3e50", alpha=alpha, linewidth=1.5)
        ax.add_patch(patch)
        
        # Label
        if obj["name"] != "lamp":
            ax.text(obj["x"], obj["z"], obj["name"], ha='center', va='center', fontsize=8, weight='bold',
                    bbox=dict(facecolor='white', alpha=0.75, edgecolor='none', pad=1.5))
                    
    # View settings
    ax.set_aspect('equal')
    ax.set_xlim(-3.2, 3.2)
    ax.set_ylim(-3.2, 3.2)
    ax.grid(True, linestyle=':', alpha=0.4)
    
    # Title & Badge
    score = metrics["score"]
    status = metrics["status"]
    
    fname = os.path.basename(file_path)
    if len(fname) > 45:
        fname = fname[:42] + "..."
    ax.set_title(f"{fname}\nScore: {score:.1f}/100 [{status}]", fontsize=9, weight='bold', color='black', pad=6)
    
    plt.tight_layout()
    os.makedirs(os.path.dirname(out_img_path), exist_ok=True)
    plt.savefig(out_img_path, dpi=150, bbox_inches='tight')
    plt.close(fig)

def build_excel_report(ranked_data, output_excel):
    """Creates a beautifully formatted Excel report with embedded layout images."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Layout Rankings Summary"
    ws.views.sheetView[0].showGridLines = True

    # Colors
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    sub_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
    
    font_title = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    # 1. Title Banner
    ws.merge_cells("A1:L2")
    cell_title = ws["A1"]
    cell_title.value = "  ROOM Layout Quality & Walkability Benchmark Report"
    cell_title.font = font_title
    cell_title.fill = header_fill
    cell_title.alignment = Alignment(vertical="center", horizontal="left")
    
    # 2. Subtitle / Summary info
    ws.merge_cells("A3:L3")
    cell_sub = ws["A3"]
    cell_sub.value = f"  Total Runs Analyzed: {len(ranked_data)}  |  Metrics: Furniture Overlaps, OBB Overlap Area, Out-of-Bounds Protrusions, Free Floor %, Walkability"
    cell_sub.font = Font(name="Segoe UI", size=10, italic=True, color="1F4E78")
    cell_sub.fill = sub_header_fill
    cell_sub.alignment = Alignment(vertical="center", horizontal="left")
    
    # 3. Column Headers
    headers = [
        "Rank", "Layout Visualization", "Log File Name", "Quality Score",
        "Overlaps (Count)", "OBB Overlap (m²)", "Out of Bounds", "Max OOB Dist (m)",
        "Free Floor %", "Walkable Components", "Center Clear (m)", "Grade Status"
    ]
    
    ws.row_dimensions[5].height = 28
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num)
        cell.value = h_text
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin

    # Column widths
    col_widths = {
        "A": 8,   "B": 24,  "C": 48,  "D": 15,  "E": 16,  "F": 18,
        "G": 15,  "H": 18,  "I": 14,  "J": 20,  "K": 16,  "L": 20
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 4. Fill Data Rows
    start_row = 6
    for idx, item in enumerate(ranked_data, 1):
        row_num = start_row + idx - 1
        ws.row_dimensions[row_num].height = 115
        
        m = item["metrics"]
        fname = os.path.basename(item["file_path"])
        img_path = item["img_path"]
        
        # Rank
        cell_rank = ws.cell(row=row_num, column=1, value=f"#{idx}")
        cell_rank.alignment = Alignment(horizontal="center", vertical="center")
        cell_rank.font = font_bold
        
        # Image
        if os.path.exists(img_path):
            img = OpenPyxlImage(img_path)
            img.width = 140
            img.height = 140
            ws.add_image(img, f"B{row_num}")
            
        # File name
        cell_file = ws.cell(row=row_num, column=3, value=fname)
        cell_file.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell_file.font = font_body
        
        # Score
        score_val = m["score"]
        cell_score = ws.cell(row=row_num, column=4, value=score_val)
        cell_score.alignment = Alignment(horizontal="center", vertical="center")
        cell_score.font = Font(name="Segoe UI", size=11, bold=True)
        cell_score.number_format = '0.0'
        
        if score_val >= 90:
            cell_score.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        elif score_val >= 75:
            cell_score.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        elif score_val >= 50:
            cell_score.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        else:
            cell_score.fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
            
        # Metrics
        ws.cell(row=row_num, column=5, value=m["furn_aabb_overlaps"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=6, value=m["total_obb_overlap_area"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=7, value=m["oob_count"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=8, value=m["max_oob_dist"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=9, value=m["free_floor_pct"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=10, value=m["walkable_components"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=11, value=m["center_clearance"]).alignment = Alignment(horizontal="center", vertical="center")
        
        # Status Grade Badge
        cell_status = ws.cell(row=row_num, column=12, value=m["status"])
        cell_status.alignment = Alignment(horizontal="center", vertical="center")
        cell_status.font = font_bold
        
        if m["status"] == "EXCELLENT":
            cell_status.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="006100")
        elif m["status"] == "GOOD":
            cell_status.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="375623")
        elif m["status"] == "MINOR ISSUES":
            cell_status.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="9C6500")
        else:
            cell_status.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            cell_status.font = Font(name="Segoe UI", size=10, bold=True, color="9C0006")
            
        # Set borders & fonts for row cells
        for c in range(1, 13):
            cell_item = ws.cell(row=row_num, column=c)
            cell_item.border = border_thin
            if c not in [4, 12]:
                cell_item.font = font_body
            if idx % 2 == 0 and c not in [4, 12]:
                cell_item.fill = zebra_fill

    wb.save(output_excel)
    print(f"Excel report successfully generated: {output_excel}")

def main():
    debug_dir = os.path.dirname(os.path.abspath(__file__))
    render_dir = os.path.join(debug_dir, "rendered_plots")
    excel_file = os.path.join(debug_dir, "layout_ranking_report.xlsx")
    
    files = sorted(glob.glob(os.path.join(debug_dir, "*.txt")), key=natural_sort_key)
    print(f"Analyzing {len(files)} log files...")
    
    scenarios = []
    for fpath in files:
        objs, overlaps = parse_log_file(fpath)
        if not objs:
            continue
        metrics = analyze_scene(objs, overlaps)
        
        fname = os.path.basename(fpath).replace(".txt", "")
        img_out = os.path.join(render_dir, f"{fname}.png")
        render_scene_image(fpath, objs, metrics, img_out)
        
        scenarios.append({
            "file_path": fpath,
            "metrics": metrics,
            "img_path": img_out
        })
        
    # Rank scenarios by Quality Score (Descending)
    scenarios.sort(key=lambda x: x["metrics"]["score"], reverse=True)
    
    print("\n" + "="*70)
    print(" TOP BEST LAYOUT CONFIGURATIONS ")
    print("="*70)
    for rank, item in enumerate(scenarios[:5], 1):
        m = item["metrics"]
        print(f"Rank #{rank}: {os.path.basename(item['file_path'])}")
        print(f"   Score: {m['score']:.1f}/100 [{m['status']}] | Overlaps: {m['furn_aabb_overlaps']} | OBB Area: {m['total_obb_overlap_area']} m² | Out-of-Bounds: {m['oob_count']}")
    print("="*70)
    
    build_excel_report(scenarios, excel_file)

if __name__ == "__main__":
    main()
